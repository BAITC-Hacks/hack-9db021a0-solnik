"""Тесты конвейера на контрольном комплекте из кейса.

Сетевой режим не требуется: сравнение запускается без агента-верификатора,
поэтому тесты проходят и без ключа OpenAI.
"""
import pytest

from src.compare import run
from src.docparse import find, parse_docx
from src.orgmap import build

BEFORE = "data/samples/polozhenie_red8.docx"
AFTER = "data/samples/polozhenie_red9.docx"


@pytest.fixture(scope="module")
def report():
    return run(BEFORE, AFTER, verify=False)


def test_clauses_keep_numbers_and_citations():
    clauses = parse_docx(BEFORE)
    assert len(clauses) > 200
    c = find(clauses, "3.4")
    assert c is not None
    assert c.cite.endswith("п. 3.4")
    assert "состоит из следующих структурных подразделений" in c.text


def test_units_before_and_after():
    """Контрольный случай: было два департамента, стало четыре."""
    before, _ = build(BEFORE)
    after, _ = build(AFTER)
    assert set(before) == {"ДНМ", "ДККМ"}
    assert set(after) == {"ДИТААД", "ДОА", "ДНМ", "ДККМ"}


def test_every_unit_has_functions_with_citations():
    after, _ = build(AFTER)
    for unit in after.values():
        assert unit.functions, f"у {unit.code} не извлечены функции"
        for fn in unit.functions:
            assert fn.cite.startswith("polozhenie_red9")
            assert fn.number


def test_created_units_detected(report):
    created = {f.title.split()[-1] for f in report.by_kind("unit_created")}
    assert created == {"ДИТААД", "ДОА"}


def test_kept_units_detected(report):
    kept = {f.title.split()[-2] for f in report.by_kind("unit_kept")}
    assert kept == {"ДНМ", "ДККМ"}


def test_every_finding_has_source_with_clause(report):
    """Главное требование ТЗ: вывод без ссылки на пункт недопустим."""
    for f in report.findings:
        assert f.sources, f"вывод без источника: {f.title}"
        assert any("п. " in s.cite for s in f.sources), f.title


def test_duplication_never_compares_clause_with_itself(report):
    for f in report.by_kind("duplication"):
        cites = [s.cite for s in f.sources]
        assert len(set(cites)) == len(cites)


def test_confidence_within_bounds(report):
    for f in report.findings:
        assert 0.0 <= f.confidence <= 1.0


def test_conflict_of_interest_found(report):
    """Контрольный случай: ДККМ разрабатывает методологию и сам оценивает качество."""
    conflicts = report.by_kind("conflict_of_interest")
    assert conflicts, "конфликт интересов не обнаружен"
    assert any("ДККМ" in f.title for f in conflicts)
    for f in conflicts:
        assert len(f.sources) == 2, "конфликт должен подтверждаться двумя пунктами"


def test_same_document_yields_no_changes():
    """Документ, сравненный сам с собой, не должен давать потерь и передач."""
    r = run(AFTER, AFTER, verify=False)
    assert not r.by_kind("function_lost")
    assert not r.by_kind("function_moved")
    assert not r.by_kind("unit_created")


def test_export_produces_valid_docx(report):
    import io
    from docx import Document
    from src.export import build_docx

    blob = build_docx(report, "Тестовое заключение.")
    assert blob[:2] == b"PK", "на выходе должен быть файл Word"
    doc = Document(io.BytesIO(blob))
    assert len(doc.tables) == 2, "таблицы структуры и отклонений"
    assert "ЗАКЛЮЧЕНИЕ" in doc.paragraphs[0].text

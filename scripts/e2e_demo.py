"""Сквозная проверка демо в настоящем браузере.

Проходит путь пользователя и жюри: стартовая страница → рабочий экран →
контрольный комплект → все разделы → поиск → работа агента → выгрузка
заключения → загрузка своих файлов → ошибочный файл → мобильная вёрстка.
Отдельно проверяет, что разметка внутри документа не выполняется в браузере.

Запуск при работающем сервере:  python scripts/e2e_demo.py
"""
from __future__ import annotations

import io
import os
import sys
import tempfile
from pathlib import Path

from docx import Document
from playwright.sync_api import sync_playwright

BASE = os.getenv("ORGTRACE_URL", "http://127.0.0.1:8000")
RED8 = str(Path("data/samples/polozhenie_red8.docx").resolve())
RED9 = str(Path("data/samples/polozhenie_red9.docx").resolve())
ANALYSIS_TIMEOUT = 240_000

results: list[tuple[bool, str, str]] = []


def check(ok: bool, name: str, detail: str = "") -> bool:
    results.append((bool(ok), name, detail))
    print(("  OK   " if ok else "  FAIL ") + name + (f" — {detail}" if detail else ""), flush=True)
    return bool(ok)


def make_xss_docx(path: str) -> None:
    """Документ, где в названии подразделения спрятана разметка со скриптом."""
    doc = Document()
    doc.add_paragraph("1. Общие положения")
    doc.add_paragraph("3. Структура")
    doc.add_paragraph('3.4. БВА состоит из следующих структурных подразделений: '
                      'а. Департамент <img src=x onerror="window.__xss=1"> проверки (ДПР).')
    doc.add_paragraph("5. Права и обязанности")
    doc.add_paragraph("5.3. Директор ДПР:")
    doc.add_paragraph('5.3.1. проводит проверки <script>window.__xss=2</script> в зоне ответственности;')
    doc.add_paragraph("5.3.2. организует работу департамента и контроль исполнения поручений;")
    doc.save(path)


def wait_analysis(page) -> None:
    page.wait_for_function(
        "() => document.querySelector('.stats') || "
        "(document.querySelector('#out .empty') && !document.querySelector('#go').disabled)",
        timeout=ANALYSIS_TIMEOUT)


def main() -> int:
    console_errors: list[str] = []
    failed_requests: list[str] = []

    with sync_playwright() as p:
        browser = p.chromium.launch()
        ctx = browser.new_context(viewport={"width": 1440, "height": 900}, locale="ru-RU",
                                  accept_downloads=True)
        page = ctx.new_page()
        page.on("pageerror", lambda e: console_errors.append(f"pageerror: {e}"))
        page.on("console", lambda m: console_errors.append(f"console: {m.text}")
                if m.type == "error" and "favicon" not in m.text and "400" not in m.text else None)
        # Скачивание файла браузер помечает как прерванный запрос (net::ERR_ABORTED),
        # хотя файл сохраняется целиком, — такие случаи сбоем не считаем.
        page.on("requestfailed", lambda r: None if ("/api/export" in r.url and "ABORTED" in (r.failure or ""))
                else failed_requests.append(f"{r.method} {r.url} ({r.failure})"))

        # ---------- стартовая страница ----------
        print("\n[1] Стартовая страница")
        r = page.goto(BASE + "/", wait_until="networkidle")
        check(r.status == 200, "главная открывается", f"HTTP {r.status}")
        check(page.title().startswith("OrgTrace"), "заголовок вкладки", page.title())
        broken = page.evaluate("""() => [...document.images]
            .filter(i => !i.complete || i.naturalWidth === 0).map(i => i.src)""")
        check(not broken, "все изображения загрузились", ", ".join(broken) or "логотип и снимок экрана")
        check(page.locator('link[rel="icon"]').count() == 1, "иконка вкладки подключена")
        for anchor in ("#problem", "#how", "#trust"):
            check(page.locator(anchor).count() == 1, f"якорь навигации {anchor} существует")
        overflow = page.evaluate("() => document.documentElement.scrollWidth > window.innerWidth")
        check(not overflow, "нет горизонтальной прокрутки на десктопе")
        links = page.evaluate("() => [...document.querySelectorAll('a[href]')].map(a => a.getAttribute('href'))")
        bad_links = []
        for href in sorted(set(links)):
            if href.startswith("#"):
                if page.locator(href).count() != 1:
                    bad_links.append(href)
            elif href.startswith("/"):
                if page.request.get(BASE + href).status != 200:
                    bad_links.append(href)
        check(not bad_links, "все ссылки главной ведут на существующие разделы и страницы",
              ", ".join(bad_links) or f"проверено ссылок: {len(set(links))}")
        check(page.locator("footer .fcol").count() == 3, "футер с тремя колонками на месте")
        check(page.evaluate("() => { const i=document.querySelector('.cta2 img'); return !!i && i.naturalWidth>0 }"),
              "иллюстрация блока «Проверьте на своём комплекте» загрузилась")
        check(page.locator("#trust .fcell").count() == 6, "сетка возможностей: 6 карточек")
        page.keyboard.press("Tab")
        focused = page.evaluate("() => document.activeElement.tagName")
        check(focused == "A", "навигация с клавиатуры: Tab попадает на ссылку", focused)
        page.locator("header .actions a.btn:not(.outline)").click()
        page.wait_for_url("**/app")
        check(page.url.endswith("/app"), "кнопка «Начать анализ» ведёт в сервис")
        page.go_back(); page.wait_for_load_state("networkidle")
        check(page.url.rstrip("/") == BASE.rstrip("/"), "кнопка браузера «Назад» возвращает на главную")
        page.go_forward(); page.wait_for_load_state("networkidle")
        check(not page.locator("#dl").is_visible(), "до анализа кнопка выгрузки скрыта")

        # ---------- рабочий экран до анализа ----------
        print("\n[2] Рабочий экран до анализа")
        page.locator('.it[data-k="duplication"]').click()
        hint = page.locator("#out .empty").inner_text()
        check("двум подразделениям" in hint, "раздел до анализа показывает подсказку")
        page.locator('.it[data-k="all"]').click()
        page.click("#go")
        msg = page.locator("#out .empty").inner_text()
        check("оба файла" in msg, "«Сравнить» без файлов объясняет, что нужно")

        # ---------- контрольный комплект ----------
        print("\n[3] Контрольный комплект")
        analyze_calls = []
        page.on("request", lambda r: analyze_calls.append(1) if r.url.endswith("/api/analyze") else None)
        page.click("#demo")
        page.evaluate("() => document.querySelector('#demo').click()")   # повторный клик во время разбора
        check(page.locator("#go").is_disabled(), "кнопки блокируются на время разбора")
        wait_analysis(page)
        check(page.locator(".stats").count() == 1, "сводка появилась")
        check(len(analyze_calls) == 1, "повторный клик во время разбора не запускает второй анализ",
              f"запросов: {len(analyze_calls)}")
        stats = page.locator(".stat b").all_inner_texts()
        check(stats and stats[0] == "4", "подразделений после: 4", str(stats))
        new_units = page.locator(".u.new").all_inner_texts()
        check(any("ДИТААД" in u for u in new_units) and any("ДОА" in u for u in new_units),
              "созданные ДИТААД и ДОА выделены", "; ".join(new_units))
        check(page.locator("#dl").is_visible(), "кнопка «Скачать заключение» появилась")
        check("п." in page.locator("pre.conc").inner_text(), "заключение ссылается на пункты")

        # ---------- все разделы сайдбара ----------
        print("\n[4] Разделы сайдбара")
        for btn in page.locator(".it[data-k]").all():
            kind = btn.get_attribute("data-k")
            label = btn.locator(".lbl").inner_text().split("\n")[0].strip()
            counter = btn.locator(".cnt").inner_text().strip()
            btn.click()
            page.wait_for_timeout(120)
            active = page.locator(f'.it[data-k="{kind}"]').get_attribute("class")
            cards = page.locator("#out .f").count()
            empty = page.locator("#out .empty").count()
            if counter.isdigit():
                ok = cards == int(counter) or (int(counter) == 0 and empty == 1)
                check(ok and "on" in active, f"«{label}»: счётчик {counter} совпадает с карточками",
                      f"карточек {cards}")
            else:
                check(page.locator("#out").inner_text().strip() != "" and "on" in active,
                      f"«{label}»: раздел открывается")

        # ---------- ссылки на пункты ----------
        print("\n[5] Ссылки на пункты во всех выводах")
        page.locator('.it[data-k="duplication"]').click()
        cites = page.locator("#out .src b").all_inner_texts()
        check(cites and all("п." in c for c in cites), "у каждого источника указан пункт",
              f"проверено ссылок: {len(cites)}")

        # ---------- поиск ----------
        print("\n[6] Поиск")
        total = page.locator("#out .f").count()
        page.fill("#q", "5.4.10")
        page.wait_for_timeout(150)
        found = page.locator("#out .f").count()
        check(0 < found < total, "поиск по номеру пункта сужает список", f"{found} из {total}")
        check(page.evaluate("() => document.activeElement.id") == "q",
              "фокус остаётся в поле поиска при вводе")
        page.fill("#q", "такого-текста-нет-нигде")
        page.wait_for_timeout(150)
        check(page.locator("#out .f").count() == 0 and "Ничего не найдено" in page.locator("#out").inner_text(),
              "пустой результат поиска объясняется")

        # ---------- работа агента ----------
        print("\n[7] Работа агента")
        agent_card = None
        for kind in ("false_positive", "function_lost"):
            page.locator(f'.it[data-k="{kind}"]').click()
            page.wait_for_timeout(120)
            cards = page.locator("#out .f").filter(has=page.locator(".vf"))
            if cards.count():
                agent_card = cards.first
                break
        if check(agent_card is not None, "есть вывод, проверенный агентом"):
            verdict = agent_card.locator(".vf").inner_text()
            check("Агент-верификатор" in verdict, "вердикт агента показан", verdict.split("\n")[0][:90])
            summary = agent_card.locator("summary")
            if summary.count():
                summary.click()
                calls = agent_card.locator(".tc").count()
                check(calls > 0, "трассировка вызовов раскрывается", f"вызовов: {calls}")

        # ---------- выгрузка ----------
        print("\n[8] Выгрузка заключения")
        with page.expect_download(timeout=30_000) as dl_info:
            page.click("#dl")
        dl = dl_info.value
        tmp = Path(tempfile.mkdtemp()) / dl.suggested_filename
        dl.save_as(tmp)
        check(dl.suggested_filename.endswith(".docx"), "файл выгрузки — .docx", dl.suggested_filename)
        try:
            d = Document(str(tmp))
            check("ЗАКЛЮЧЕНИЕ" in d.paragraphs[0].text, "заключение открывается в Word",
                  f"таблиц {len(d.tables)}, строк отклонений {len(d.tables[-1].rows) - 1}")
        except Exception as e:
            check(False, "заключение открывается в Word", str(e))

        # ---------- свои файлы ----------
        print("\n[9] Загрузка своих файлов")
        page.set_input_files("#before", RED8)
        page.set_input_files("#after", RED9)
        page.click("#go")
        wait_analysis(page)
        mode_line = page.locator(".mode").first.inner_text() if page.locator(".mode").count() else ""
        check("polozhenie_red8.docx" in mode_line and "Temp" not in mode_line,
              "в результатах настоящие имена файлов", mode_line[:100])

        # ---------- ошибочный файл ----------
        print("\n[10] Ошибочный файл")
        bad = Path(tempfile.mkdtemp()) / "not_word.docx"
        bad.write_bytes(b"this is not a word document")
        page.set_input_files("#before", str(bad))
        page.set_input_files("#after", RED9)
        page.click("#go")
        page.wait_for_function("() => !document.querySelector('#go').disabled", timeout=30_000)
        err = page.locator("#out .empty").inner_text()
        check("повреждён" in err or "не является" in err, "понятное сообщение об ошибке", err[:90])

        # ---------- внедрение разметки через документ ----------
        print("\n[11] Разметка внутри документа не выполняется")
        xss = Path(tempfile.mkdtemp()) / "xss.docx"
        make_xss_docx(str(xss))
        page.evaluate("() => { window.__xss = 0; }")
        page.set_input_files("#before", str(xss))
        page.set_input_files("#after", str(xss))
        page.click("#go")
        wait_analysis(page)
        page.locator('.it[data-k="units"]').click()
        page.wait_for_timeout(300)
        check(page.evaluate("() => window.__xss") == 0, "скрипт из документа не выполнился")
        check("<img" in page.locator("#out").inner_text(),
              "разметка показана как текст", "виден литерал <img …>")
        check(page.locator("#out img").count() == 0, "в разметку не попало ни одного <img>")

        # ---------- мобильная вёрстка ----------
        print("\n[12] Мобильная вёрстка")
        mob = ctx.new_page()
        mob.set_viewport_size({"width": 390, "height": 844})
        mob.goto(BASE + "/", wait_until="networkidle")
        check(not mob.evaluate("() => document.documentElement.scrollWidth > window.innerWidth"),
              "главная без горизонтальной прокрутки на телефоне")
        mob.goto(BASE + "/app", wait_until="networkidle")
        check(not mob.locator(".side").is_visible() or
              mob.evaluate("() => document.querySelector('.side').getBoundingClientRect().right <= 0"),
              "сайдбар скрыт на телефоне")
        mob.click("#burger")
        mob.wait_for_timeout(300)
        check(mob.evaluate("() => document.querySelector('.side').getBoundingClientRect().left >= 0"),
              "бургер открывает сайдбар")
        mob.click("#scrim", position={"x": 370, "y": 400})
        mob.wait_for_timeout(300)
        check(mob.evaluate("() => document.querySelector('.side').getBoundingClientRect().right <= 0"),
              "тап мимо закрывает сайдбар")
        check(not mob.evaluate("() => document.documentElement.scrollWidth > window.innerWidth"),
              "рабочий экран без горизонтальной прокрутки на телефоне")

        # ---------- возврат ----------
        page.click('a.it[href="/"]')
        page.wait_for_load_state("networkidle")
        check(page.url.rstrip("/") == BASE.rstrip("/"), "«На главную» возвращает на главную")

        browser.close()

    print("\n[13] Браузер")
    check(not console_errors, "нет ошибок JavaScript", "; ".join(console_errors[:3]))
    check(not failed_requests, "нет сорвавшихся запросов", "; ".join(failed_requests[:3]))

    passed = sum(ok for ok, _, _ in results)
    print(f"\nИТОГО: {passed} из {len(results)} проверок пройдено")
    for ok, name, detail in results:
        if not ok:
            print(f"  не пройдено: {name} — {detail}")
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())

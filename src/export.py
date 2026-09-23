"""Выгрузка итогового заключения в Word.

ТЗ, раздел «Артефакты»: итоговое аналитическое заключение. На экране его видно,
но аналитик уносит результат руководителю файлом, поэтому заключение собирается
в .docx — с таблицей выводов и ссылками на пункты исходных документов.
"""
from __future__ import annotations

import io
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from .compare import Report

KIND_RU = {
    "unit_created": "Создано подразделение",
    "unit_removed": "Подразделение отсутствует",
    "unit_kept": "Подразделение сохранено",
    "function_lost": "Возможная потеря функции",
    "false_positive": "Снято агентом",
    "function_moved": "Функция передана",
    "function_generalized": "Функция стала общей",
    "duplication": "Дублирование",
    "conflict_of_interest": "Конфликт интересов",
}
SEV_RU = {"high": "высокий", "medium": "средний", "info": "справочно"}

# В файл идут выводы, требующие решения; справочные остаются в сервисе.
REPORTABLE = ("conflict_of_interest", "function_lost", "duplication",
              "function_generalized", "function_moved", "unit_created", "unit_removed")


def _heading(doc: Document, text: str, size: int = 13) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    p.space_after = Pt(6)


def build_docx(report: Report, conclusion: str) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ЗАКЛЮЧЕНИЕ\nпо результатам анализа организационной структуры и функционала")
    run.bold = True
    run.font.size = Pt(14)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    m = meta.add_run(f"Комплект «до»: {report.before_doc}    Комплект «после»: {report.after_doc}\n"
                     f"Дата анализа: {date.today().strftime('%d.%m.%Y')}    Сформировано: OrgTrace")
    m.font.size = Pt(9)
    m.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    _heading(doc, "1. Аналитическое заключение")
    for block in conclusion.split("\n"):
        if block.strip():
            doc.add_paragraph(block.strip())

    _heading(doc, "2. Структура подразделений")
    before_codes = {u["code"] for u in report.units_before}
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, head in enumerate(["Код", "Наименование", "Статус", "Функций"]):
        cell = table.rows[0].cells[i]
        cell.text = head
        cell.paragraphs[0].runs[0].bold = True
    for u in report.units_after:
        row = table.add_row().cells
        row[0].text = u["code"]
        row[1].text = u["name"]
        row[2].text = "сохранено" if u["code"] in before_codes else "создано"
        row[3].text = str(len(u["functions"]))

    _heading(doc, "3. Выявленные отклонения")
    findings = [f for f in report.findings if f.kind in REPORTABLE]
    if not findings:
        doc.add_paragraph("Отклонений, требующих решения, не выявлено.")
    else:
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for i, head in enumerate(["№", "Тип", "Вывод", "Основание (пункт)", "Риск"]):
            cell = table.rows[0].cells[i]
            cell.text = head
            cell.paragraphs[0].runs[0].bold = True
        for n, f in enumerate(findings, 1):
            row = table.add_row().cells
            row[0].text = str(n)
            row[1].text = KIND_RU.get(f.kind, f.kind)
            detail = f.detail
            if f.verification:
                detail += f" Проверка агентом: {f.verification.get('verdict_ru', '')}."
            if f.recommendation:
                detail += f" Рекомендация: {f.recommendation}"
            row[2].text = f"{f.title}. {detail}"
            row[3].text = "\n".join(s.cite for s in f.sources)
            row[4].text = SEV_RU.get(f.severity, f.severity)

    _heading(doc, "4. Порядок формирования выводов")
    doc.add_paragraph(
        "Состав подразделений и принадлежность функций извлечены из структуры документов "
        "программно; номера пунктов взяты из исходных файлов и не формируются моделью. "
        "Сопоставление формулировок выполнено по смысловой близости "
        f"(режим: {report.mode}). Каждое подозрение на потерю функции дополнительно "
        "перепроверено агентом, который искал подтверждение в комплекте «после» и обязан "
        "был указать пункт; вывод без ссылки на пункт не принимается.")
    note = doc.add_paragraph()
    r = note.add_run("Выводы носят рекомендательный характер и требуют проверки "
                     "ответственным сотрудником.")
    r.italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
